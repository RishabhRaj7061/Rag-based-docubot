
import io
import logging
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from docx import Document as WordDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from backend.services.document_service import (
    UPLOAD_DIR,
    INDEX_DIR,
    load_documents,
    save_documents,
    process_pdf_file,
)
from backend.services.pdf_service import chunk_pages
from backend.services.vector_service import save_document_vectors
from backend.services.qa_service import ask_document

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("docubot")

app = FastAPI(
    title="DocuBot API",
    description="RAG-based document question-answering API",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
INDEX_DIR.mkdir(parents=True, exist_ok=True)

MAX_FILE_SIZE = 200 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc"}


class AskRequest(BaseModel):
    question: str = Field(..., min_length=1, max_length=10000)
    document_id: str = Field(..., min_length=1)
    top_k: int = Field(default=5, ge=1, le=20)


@app.get("/")
def root():
    return {
        "message": "DocuBot API is running",
        "docs": "/docs",
        "health": "/health",
    }


@app.get("/health")
def health():
    return {"status": "ok", "service": "DocuBot API"}


@app.get("/documents")
def list_documents():
    """Return saved document metadata for the Streamlit frontend."""
    return load_documents()


def convert_doc_to_docx(content: bytes, filename: str) -> bytes:
    """
    Convert legacy .doc files to .docx using LibreOffice.
    LibreOffice must be installed for legacy .doc uploads.
    """
    libreoffice = shutil.which("soffice") or shutil.which("libreoffice")

    if not libreoffice:
        raise HTTPException(
            status_code=400,
            detail=(
                "Legacy .doc files require LibreOffice for conversion. "
                "Please save the file as .docx and upload it again."
            ),
        )

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        source_path = temp_path / Path(filename).name
        source_path.write_bytes(content)

        try:
            subprocess.run(
                [
                    libreoffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(temp_path),
                    str(source_path),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=60,
            )
        except subprocess.TimeoutExpired:
            raise HTTPException(
                status_code=400,
                detail="Word document conversion timed out.",
            )
        except subprocess.CalledProcessError as exc:
            logger.error("LibreOffice conversion failed: %s", exc.stderr)
            raise HTTPException(
                status_code=400,
                detail="Could not convert the .doc file. Try saving it as .docx.",
            )

        converted_path = temp_path / f"{source_path.stem}.docx"

        if not converted_path.exists():
            raise HTTPException(
                status_code=400,
                detail="Word conversion did not produce a .docx file.",
            )

        return converted_path.read_bytes()


def process_word_file(filename: str, content: bytes) -> dict:
    """Extract DOCX text, create chunks and save a vector index."""
    document_id = str(uuid.uuid4())
    safe_filename = Path(filename).name
    stored_path = UPLOAD_DIR / f"{document_id}.docx"

    try:
        word_doc = WordDocument(io.BytesIO(content))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in word_doc.paragraphs
            if paragraph.text and paragraph.text.strip()
        ]

        # Include text from tables, which is common in resumes.
        for table in word_doc.tables:
            for row in table.rows:
                cell_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text and cell.text.strip()
                )
                if cell_text:
                    paragraphs.append(cell_text)

        full_text = "\n".join(paragraphs).strip()

        if not full_text:
            raise ValueError(
                "No readable text was found in this Word document."
            )

        pages = [{"page": 1, "text": full_text}]
        chunks = chunk_pages(pages)

        if not chunks:
            raise ValueError("Could not create text chunks from this document.")

        # Store consistent metadata for the document and its index.
        for chunk in chunks:
            chunk["document_id"] = document_id
            chunk["filename"] = safe_filename
            chunk["page_number"] = chunk.get("page", 1)

        vector_info = save_document_vectors(document_id, chunks)

        # Save the original content after successful indexing.
        stored_path.write_bytes(content)

        record = {
            "document_id": document_id,
            "filename": safe_filename,
            "stored_filename": stored_path.name,
            "page_count": 1,
            "chunk_count": len(chunks),
            "embedding_dimensions": vector_info.get(
                "embedding_dimensions", 384
            ),
            "status": "indexed",
        }

        documents = load_documents()
        documents.append(record)
        save_documents(documents)

        return {
            **record,
            "pages_with_text": 1,
            "vectors_created": vector_info.get(
                "vectors_created", len(chunks)
            ),
        }

    except Exception:
        stored_path.unlink(missing_ok=True)

        # Remove only the files/folder belonging to this new document.
        index_folder = INDEX_DIR / document_id
        if index_folder.is_dir():
            shutil.rmtree(index_folder, ignore_errors=True)

        (INDEX_DIR / f"{document_id}.faiss").unlink(missing_ok=True)
        (INDEX_DIR / f"{document_id}_chunks.json").unlink(missing_ok=True)
        raise


@app.post("/documents/upload")
async def upload_document(file: UploadFile = File(...)):
    """
    Upload a PDF or Word document.

    IMPORTANT: The multipart field name is 'file', matching the frontend:
    files={'file': (filename, content, content_type)}
    """
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename was provided.")

    filename = Path(file.filename).name
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Upload a PDF, DOCX or DOC file.",
        )

    try:
        content = await file.read(MAX_FILE_SIZE + 1)

        if not content:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty.",
            )

        if len(content) > MAX_FILE_SIZE:
            raise HTTPException(
                status_code=413,
                detail="The file exceeds the 200 MB upload limit.",
            )

        if extension == ".pdf":
            result = process_pdf_file(filename, content)

        else:
            if extension == ".doc":
                content = convert_doc_to_docx(content, filename)

            result = process_word_file(filename, content)

        logger.info(
            "Document indexed: %s (%s)",
            result.get("filename"),
            result.get("document_id"),
        )

        return {
            "message": "Document uploaded and indexed successfully.",
            **result,
        }

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Document upload/indexing failed: %s", filename)
        raise HTTPException(
            status_code=500,
            detail=f"Could not process '{filename}': {str(exc)}",
        )
    finally:
        await file.close()


@app.post("/ask")
def ask_question(request: AskRequest):
    """Search a document's vector index and generate an answer."""
    question = request.question.strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Please enter a question.",
        )

    documents = load_documents()
    document = next(
        (
            item
            for item in documents
            if str(item.get("document_id", "")) == request.document_id
        ),
        None,
    )

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found. Refresh the document list and try again.",
        )

    try:
        result = ask_document(
            question=question,
            document_id=request.document_id,
            top_k=request.top_k,
        )

        if isinstance(result, str):
            result = {"answer": result, "sources": []}

        if not isinstance(result, dict):
            raise ValueError("The question-answering service returned an invalid response.")

        return {
            "answer": result.get("answer", "No answer was generated."),
            "sources": result.get("sources", []),
            "document_id": request.document_id,
            "filename": document.get("filename", "Unknown document"),
        }

    except FileNotFoundError as exc:
        logger.exception("Index not found for document %s", request.document_id)
        raise HTTPException(
            status_code=404,
            detail=(
                "The search index for this document is missing or uses an "
                "older format. Upload the document again to rebuild its index."
            ),
        ) from exc

    except HTTPException:
        raise

    except Exception as exc:
        logger.exception("Question answering failed")
        raise HTTPException(
            status_code=502,
            detail=(
                "Question answering failed. Check the backend terminal for "
                f"the underlying error. Details: {str(exc)}"
            ),
        ) from exc
